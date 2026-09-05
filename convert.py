import os
import re
import base64
import urllib.request
import urllib.parse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# XML helpers for styling docx
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_paragraph_runs(p, text):
    # Regex parser for bold (**text**), italic (*text*), and inline code (`code`)
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            content = part[2:-2]
            r = p.add_run(content)
            r.bold = True
        elif part.startswith('*') and part.endswith('*'):
            content = part[1:-1]
            r = p.add_run(content)
            r.italic = True
        elif part.startswith('`') and part.endswith('`'):
            content = part[1:-1]
            r = p.add_run(content)
            r.font.name = 'Courier New'
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(160, 30, 30)
        else:
            p.add_run(part)

def create_docx_table(doc, table_data):
    if not table_data:
        return
    rows = len(table_data)
    cols = len(table_data[0])
    table = doc.add_table(rows=rows, cols=cols)
    table.autofit = True
    
    for r_idx, row in enumerate(table_data):
        for c_idx, val in enumerate(row):
            if c_idx >= cols:
                continue
            cell = table.cell(r_idx, c_idx)
            cell.text = "" # Clear default
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            add_paragraph_runs(p, val)
            
            # Header styling
            if r_idx == 0:
                set_cell_background(cell, "2D2D30")
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                # Zebra striping
                if r_idx % 2 == 0:
                    set_cell_background(cell, "F5F5F5")
            
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            
            # Border styles
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = OxmlElement('w:tcBorders')
            for border_name in ['top', 'left', 'bottom', 'right']:
                b = OxmlElement(f'w:{border_name}')
                b.set(qn('w:val'), 'single')
                b.set(qn('w:sz'), '4') # Thin border
                b.set(qn('w:color'), 'D3D3D3')
                tcBorders.append(b)
            tcPr.append(tcBorders)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def download_mermaid_image(mermaid_code, filename):
    try:
        # Base64 urlsafe encode
        graph_bytes = mermaid_code.encode("utf-8")
        base64_bytes = base64.urlsafe_b64encode(graph_bytes)
        base64_string = base64_bytes.decode("ascii")
        url = f"https://mermaid.ink/img/{base64_string}"
        
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=12) as response:
            with open(filename, 'wb') as out_file:
                out_file.write(response.read())
        return True
    except Exception as e:
        print(f"Skipping Mermaid PNG download due to network/API: {e}")
        return False

def parse_md_to_docx(md_path, docx_path):
    doc = Document()
    
    # Page Setup - Standard 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles Setup
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(30, 30, 30)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_block_lang = ""
    code_block_lines = []
    
    in_table = False
    table_data = []
    
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Handle Code Blocks
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                
                # Process code block
                if code_block_lang == "mermaid":
                    mermaid_code = "\n".join(code_block_lines)
                    temp_img = "temp_mermaid.png"
                    if download_mermaid_image(mermaid_code, temp_img):
                        p = doc.add_paragraph()
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p.paragraph_format.space_before = Pt(6)
                        p.paragraph_format.space_after = Pt(6)
                        p.add_run().add_picture(temp_img, width=Inches(5.0))
                        
                        # Add a caption in a separate paragraph styled as 'Caption'
                        p_cap = doc.add_paragraph(style='Caption')
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_cap.paragraph_format.space_before = Pt(2)
                        p_cap.paragraph_format.space_after = Pt(6)
                        caption = p_cap.add_run("Figure: System Process Diagram")
                        caption.italic = True
                        caption.font.size = Pt(9.5)
                        caption.font.color.rgb = RGBColor(100, 100, 100)
                        
                        if os.path.exists(temp_img):
                            os.remove(temp_img)
                    else:
                        p = doc.add_paragraph()
                        p.paragraph_format.left_indent = Inches(0.25)
                        p.paragraph_format.space_before = Pt(6)
                        p.paragraph_format.space_after = Pt(6)
                        r = p.add_run(mermaid_code)
                        r.font.name = 'Courier New'
                        r.font.size = Pt(9)
                        r.font.color.rgb = RGBColor(50, 50, 50)
                else:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Inches(0.25)
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(6)
                    code_text = "\n".join(code_block_lines)
                    r = p.add_run(code_text)
                    r.font.name = 'Courier New'
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = RGBColor(80, 80, 80)
                
                code_block_lines = []
            else:
                in_code_block = True
                code_block_lang = stripped[3:].strip().lower()
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line.rstrip('\n'))
            i += 1
            continue

        # Handle Tables
        if stripped.startswith("|"):
            in_table = True
            cells = [c.strip() for c in line.split("|")[1:-1]]
            if all(re.match(r'^:?-+:?$', c) for c in cells):
                i += 1
                continue
            table_data.append(cells)
            i += 1
            continue
        else:
            if in_table:
                create_docx_table(doc, table_data)
                table_data = []
                in_table = False

        if not stripped or stripped.startswith("<!--"):
            i += 1
            continue

        # Handle Images
        if stripped.startswith("![") and stripped.endswith(")"):
            match = re.match(r'^!\[(.*?)\]\((.*?)\)', stripped)
            if match:
                alt_text = match.group(1)
                img_path = match.group(2)
                
                # Resolve relative path to workspace root
                full_img_path = img_path
                if not os.path.isabs(full_img_path):
                    full_img_path = os.path.join(r"e:\VS\WPFImageTransfer", img_path)
                    
                if os.path.exists(full_img_path):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(6)
                    p.paragraph_format.space_after = Pt(6)
                    p.add_run().add_picture(full_img_path, width=Inches(5.0))
                    
                    # Add a caption in a separate paragraph styled as 'Caption'
                    p_cap = doc.add_paragraph(style='Caption')
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_before = Pt(2)
                    p_cap.paragraph_format.space_after = Pt(6)
                    caption = p_cap.add_run(f"Figure: {alt_text}")
                    caption.italic = True
                    caption.font.size = Pt(9.5)
                    caption.font.color.rgb = RGBColor(100, 100, 100)
                else:
                    p = doc.add_paragraph()
                    p.add_run(f"[Image Missing: {img_path}]").italic = True
                i += 1
                continue

        # Headers
        if stripped.startswith("# "):
            title = stripped[2:]
            p = doc.add_heading(level=0)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(title)
            r.font.name = 'Calibri'
            r.bold = True
            r.font.color.rgb = RGBColor(0, 102, 204)
        elif stripped.startswith("## "):
            title = stripped[3:]
            p = doc.add_heading(level=1)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(title)
            r.font.name = 'Calibri'
            r.bold = True
            r.font.color.rgb = RGBColor(0, 102, 204)
            
            if title.strip().lower() == "list of figures":
                p_fld = doc.add_paragraph()
                p_fld.paragraph_format.space_before = Pt(6)
                p_fld.paragraph_format.space_after = Pt(6)
                
                fldSimple = OxmlElement('w:fldSimple')
                fldSimple.set(qn('w:instr'), 'TOC \\t "Caption" \\h')
                
                run = OxmlElement('w:r')
                rPr = OxmlElement('w:rPr')
                
                color = OxmlElement('w:color')
                color.set(qn('w:val'), '808080')
                rPr.append(color)
                
                italic = OxmlElement('w:i')
                rPr.append(italic)
                
                run.append(rPr)
                text = OxmlElement('w:t')
                text.text = "[Right-click here and select 'Update Field' to update the List of Figures]"
                run.append(text)
                fldSimple.append(run)
                
                p_fld._p.append(fldSimple)
        elif stripped.startswith("### "):
            title = stripped[4:]
            p = doc.add_heading(level=2)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(title)
            r.font.name = 'Calibri'
            r.bold = True
            r.font.color.rgb = RGBColor(50, 50, 50)
        elif stripped.startswith("#### "):
            title = stripped[5:]
            p = doc.add_heading(level=3)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(title)
            r.font.name = 'Calibri'
            r.bold = True
            r.italic = True
            r.font.color.rgb = RGBColor(80, 80, 80)
            
        # Lists
        elif stripped.startswith("* ") or stripped.startswith("- "):
            content = stripped[2:]
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            add_paragraph_runs(p, content)
        elif re.match(r'^\d+\.\s', stripped):
            match = re.match(r'^(\d+)\.\s(.*)', stripped)
            content = match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            add_paragraph_runs(p, content)
            
        # Standard Paragraphs
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            add_paragraph_runs(p, stripped)

        i += 1

    # Catch if table was at the end of the file
    if in_table and table_data:
        create_docx_table(doc, table_data)

    doc.save(docx_path)
    print(f"Successfully converted MD to DOCX: {docx_path}")

if __name__ == "__main__":
    src = r"e:\VS\WPFImageTransfer\ProjectReport.md"
    dest = r"e:\VS\WPFImageTransfer\ProjectReport.docx"
    parse_md_to_docx(src, dest)
